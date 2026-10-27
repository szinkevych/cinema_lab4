"""Розділ «Звіти»: виручка по залах за період (процедура), по фільмах (представлення), по днях."""
from io import BytesIO
from datetime import timedelta
from pathlib import Path

from flask import Blueprint, abort, redirect, render_template, request, send_file, url_for

from auth import roles_required
from database import query_all, query_one, transaction
from validators import get_date

bp = Blueprint("reports", __name__, url_prefix="/reports")


def spreadsheet_text(value):
    text = str(value)
    if text.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


@bp.route("/")
@roles_required("admin")
def reports_page():
    report = load_report()
    return render_template("reports.html", **report)


@bp.route("/export.xlsx")
@roles_required("admin")
def export_excel():
    report = load_report()
    if report["errors"]:
        return redirect(url_for("reports.reports_page", **request.args))

    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Огляд"
    summary.append(["Звіт кінотеатру"])
    summary.append(["Період", f"{report['date_from']} — {report['date_to']}"])
    summary.append([])
    summary.append(["Виручка за залами", report["halls_total"]])
    summary.append(["Фільмів у звіті", len(report["movies"])])
    summary.append(["Днів із продажами", len(report["days"])])

    halls_sheet = workbook.create_sheet("Зали")
    halls_sheet.append(["Зал", "Тип", "Сеансів", "Продано квитків", "Виручка, грн"])
    for hall in report["halls"]:
        halls_sheet.append([
            spreadsheet_text(hall["hall_name"]), spreadsheet_text(hall["hall_kind"]), hall["screenings_count"],
            hall["tickets_sold"], float(hall["revenue"]),
        ])
    if report["halls"]:
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "Виручка за залами"
        chart.y_axis.title = "Зал"
        chart.x_axis.title = "Виручка, грн"
        chart.add_data(Reference(halls_sheet, min_col=5, min_row=1, max_row=len(report["halls"]) + 1),
                       titles_from_data=True)
        chart.set_categories(Reference(halls_sheet, min_col=1, min_row=2, max_row=len(report["halls"]) + 1))
        halls_sheet.add_chart(chart, "G2")

    movies_sheet = workbook.create_sheet("Фільми")
    movies_sheet.append(["Фільм", "Продано квитків", "Виручка, грн"])
    for movie in report["movies"]:
        movies_sheet.append([
            spreadsheet_text(movie["title"]), movie["tickets_sold"], float(movie["revenue"] or 0)
        ])

    days_sheet = workbook.create_sheet("По днях")
    days_sheet.append(["Дата сеансу", "Квитків", "Виручка, грн", "Наростаючим підсумком, грн"])
    for day in report["days"]:
        days_sheet.append([
            day["day"], day["tickets"], float(day["revenue"] or 0),
            float(day["running_total"] or 0),
        ])
    if report["days"]:
        chart = LineChart()
        chart.style = 13
        chart.title = "Виручка по днях"
        chart.y_axis.title = "Виручка, грн"
        chart.x_axis.title = "Дата сеансу"
        chart.add_data(Reference(days_sheet, min_col=3, min_row=1, max_row=len(report["days"]) + 1),
                       titles_from_data=True)
        chart.set_categories(Reference(days_sheet, min_col=1, min_row=2, max_row=len(report["days"]) + 1))
        days_sheet.add_chart(chart, "F2")

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17283A")
            cell.alignment = Alignment(vertical="center")
        for column in sheet.columns:
            column_letter = get_column_letter(column[0].column)
            sheet.column_dimensions[column_letter].width = min(
                max(max(len(str(cell.value or "")) for cell in column) + 2, 14), 42
            )
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    heading = str(sheet.cell(row=1, column=cell.column).value or "")
                    is_currency = "виручк" in heading.lower() or "підсумком" in heading.lower()
                    if sheet.title == "Огляд" and cell.row == 4 and cell.column == 2:
                        is_currency = True
                    cell.number_format = '#,##0.00' if is_currency else '#,##0'
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 34

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"cinema_report_{report['date_from']}_{report['date_to']}.xlsx",
    )


@bp.route("/export.pdf")
@roles_required("admin")
def export_pdf():
    report = load_report()
    if report["errors"]:
        return redirect(url_for("reports.reports_page", **request.args))

    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    font_candidates = (
        (Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/arialbd.ttf")),
        (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
         Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")),
        (Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
         Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf")),
        (Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
         Path("/System/Library/Fonts/Supplemental/Arial Bold.ttf")),
    )
    fonts = next(
        ((regular, bold) for regular, bold in font_candidates if regular.is_file() and bold.is_file()),
        None,
    )
    if fonts is None:
        abort(503, description="Не знайдено системний шрифт із підтримкою кирилиці для PDF.")
    pdfmetrics.registerFont(TTFont("CinemaReport", str(fonts[0])))
    pdfmetrics.registerFont(TTFont("CinemaReport-Bold", str(fonts[1])))

    navy = colors.HexColor("#17283A")
    orange = colors.HexColor("#F15A3B")
    pale = colors.HexColor("#EEF2F5")
    muted = colors.HexColor("#647386")
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CinemaTitle", parent=styles["Title"], fontName="CinemaReport-Bold",
        fontSize=22, leading=27, textColor=navy, alignment=TA_LEFT, spaceAfter=4 * mm,
    )
    section_style = ParagraphStyle(
        "CinemaSection", parent=styles["Heading2"], fontName="CinemaReport-Bold",
        fontSize=13, leading=16, textColor=navy, spaceBefore=5 * mm, spaceAfter=2 * mm,
    )
    body_style = ParagraphStyle(
        "CinemaBody", parent=styles["BodyText"], fontName="CinemaReport",
        fontSize=8.5, leading=11, textColor=navy,
    )
    header_style = ParagraphStyle(
        "CinemaHeader", parent=body_style, fontName="CinemaReport-Bold",
        fontSize=8, leading=10, textColor=colors.white,
    )
    note_style = ParagraphStyle(
        "CinemaNote", parent=body_style, fontSize=8, textColor=muted,
    )
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer, pagesize=landscape(A4), rightMargin=15 * mm, leftMargin=15 * mm,
        topMargin=16 * mm, bottomMargin=17 * mm,
        title=f"Звіт кінотеатру {report['date_from']} — {report['date_to']}",
        author="Кінокадр",
    )
    story = [
        Paragraph("Звіт кінотеатру «Кінокадр»", title_style),
        Paragraph(f"Період виручки за залами: {report['date_from']} — {report['date_to']}", note_style),
        Spacer(1, 4 * mm),
    ]

    def add_table(title, headers, rows, widths):
        story.append(Paragraph(title, section_style))
        data = [[Paragraph(str(value), header_style) for value in headers]]
        data.extend([
            [Paragraph(str(value), body_style) if isinstance(value, str) else value for value in row]
            for row in rows
        ])
        table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), navy),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, pale]),
            ("LINEBELOW", (0, 0), (-1, 0), 1.2, orange),
            ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#DCE2E8")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 7),
            ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.append(table)

    story.append(Paragraph("Ключові показники", section_style))
    summary_rows = [
        ["Виручка за обраний період по залах", f"{report['halls_total']:,.2f} грн"],
        ["Фільмів у звіті", str(len(report["movies"]))],
        ["Днів із продажами", str(len(report["days"]))],
    ]
    summary_data = [[Paragraph(label, body_style), Paragraph(value, body_style)]
                    for label, value in summary_rows]
    summary_table = Table(summary_data, colWidths=[95 * mm, 55 * mm], hAlign="LEFT")
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), pale),
        ("TEXTCOLOR", (1, 0), (1, -1), orange),
        ("FONTNAME", (1, 0), (1, -1), "CinemaReport-Bold"),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#DCE2E8")),
        ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.white),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    story.extend([summary_table, Spacer(1, 3 * mm)])

    add_table(
        "Виручка по залах",
        ["Зал", "Тип", "Сеансів", "Продано квитків", "Виручка, грн"],
        [[
            row["hall_name"], row["hall_kind"], row["screenings_count"],
            row["tickets_sold"], f"{row['revenue']:,.2f}",
        ] for row in report["halls"]] or [["За цей період даних немає", "", "", "", ""]],
        [52 * mm, 35 * mm, 35 * mm, 45 * mm, 48 * mm],
    )
    add_table(
        "Виручка по фільмах",
        ["Фільм", "Продано квитків", "Виручка, грн"],
        [[row["title"], row["tickets_sold"], f"{row['revenue'] or 0:,.2f}"]
         for row in report["movies"]] or [["Даних немає", "", ""]],
        [110 * mm, 52 * mm, 53 * mm],
    )
    add_table(
        "Виручка по днях",
        ["Дата сеансу", "Квитків", "Виручка, грн", "Наростаючим підсумком, грн"],
        [[
            row["day"].strftime("%d.%m.%Y"), row["tickets"],
            f"{row['revenue'] or 0:,.2f}", f"{row['running_total'] or 0:,.2f}",
        ] for row in report["days"]] or [["Даних немає", "", "", ""]],
        [45 * mm, 35 * mm, 55 * mm, 70 * mm],
    )

    def draw_page(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#DCE2E8"))
        canvas.line(doc.leftMargin, 12 * mm, landscape(A4)[0] - doc.rightMargin, 12 * mm)
        canvas.setFont("CinemaReport", 8)
        canvas.setFillColor(muted)
        canvas.drawString(doc.leftMargin, 7 * mm, "Кінотеатр «Кінокадр» · Звіт сформовано автоматично")
        canvas.drawRightString(landscape(A4)[0] - doc.rightMargin, 7 * mm, f"Сторінка {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=draw_page, onLaterPages=draw_page)
    buffer.seek(0)
    return send_file(
        buffer,
        mimetype="application/pdf",
        as_attachment=True,
        download_name=f"cinema_report_{report['date_from']}_{report['date_to']}.pdf",
    )


def load_report():
    today = query_one("SELECT CURRENT_DATE AS d")["d"]
    errors = {}
    date_from = get_date(request.args, "from", "Початок періоду", errors, required=False)
    date_to = get_date(request.args, "to", "Кінець періоду", errors, required=False)
    if "from" not in errors and date_from is None:
        date_from = today - timedelta(days=7)
    if "to" not in errors and date_to is None:
        date_to = today + timedelta(days=7)
    if not errors and date_from > date_to:
        errors["to"] = "Кінець періоду не може бути раніше за початок."

    halls, halls_total = [], 0
    if not errors:
        with transaction() as cur:
            cur.execute("CALL hall_revenue_report_proc(%s, %s, NULL)", (date_from, date_to))
            halls = cur.fetchone()["p_report"]
        halls_total = sum(h["revenue"] for h in halls)
        top = max((h["revenue"] for h in halls), default=0) or 1
        for h in halls:
            h["share"] = round(float(h["revenue"]) / float(top) * 100)

    movies = query_all("SELECT title, tickets_sold, revenue FROM movie_revenue ORDER BY revenue DESC, title")
    days = query_all("""
        SELECT s.start_time::date AS day, COUNT(t.ticket_id) AS tickets, SUM(t.price) AS revenue,
               SUM(SUM(t.price)) OVER (ORDER BY s.start_time::date) AS running_total
        FROM tickets t
        JOIN screenings s ON s.screening_id = t.screening_id
        WHERE t.status = 'sold'
        GROUP BY s.start_time::date
        ORDER BY day""")
    return {
        "halls": halls, "halls_total": halls_total, "movies": movies,
        "days": days, "errors": errors, "date_from": date_from, "date_to": date_to,
        "form_values": {
            "from": request.args.get("from", date_from.isoformat() if date_from else ""),
            "to": request.args.get("to", date_to.isoformat() if date_to else ""),
        },
    }